"""Build Supplementary Data S3-2 from the collaborator's Prism file.
"""

import argparse
import csv
import io
import json
import math
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO = Path(__file__).resolve().parent.parent

PRISM = Path(
    "/Users/henrysanford/dev/test_data/texh_revisions/nucleotide analysis_CW.prism"
)
OUT = REPO / "supp_data" / "Data S3-2.xlsx"

# The three slots every supplementary sheet fills, kept module-level so the wording lives in
# one place - the `S2_6_SHEET`/`S2_6_TITLE`/`S2_6_DESCRIPTION` pattern in `solubility.py`.
# `S3_2_SHEET` is the tab name (Excel caps a tab name at 31 characters) and the `contents`
# title column; `S3_2_TITLE` is row 1 and, minus its `S3-2 ` prefix, the `contents`
# description column; `S3_2_DESCRIPTION` is row 2.
S3_2_SHEET = "S3-2 Metabolite ratios"
# No trailing `(related to ...)` clause: unlike its neighbours this sheet accompanies no
# figure - it is a standalone table the manuscript calls out in text.
S3_2_TITLE = (
    "S3-2 Nucleotide and glutathione ratios in activated, acutely, and chronically "
    "stimulated T cells"
)
S3_2_DESCRIPTION = (
    "Ratios of nucleoside to nucleotide, nucleoside triphosphate to diphosphate, and "
    "oxidized to reduced glutathione in human T cells activated for 2 days (D2) or cultured "
    "for 8 days with (D8C) or without (D8A) chronic stimulation. Each ratio is normalized to "
    "its D2 value within a donor, so every D2 value is 1 by construction. P-values are from "
    "a Friedman test with Dunn's multiple comparisons test (adjusted p-values), computed in "
    "GraphPad Prism 10. Friedman is a repeated-measures test, so a donor missing any one "
    "condition is excluded from that ratio entirely and is not listed: GTP:GDP is from n = 3 "
    "donors, every other ratio from n = 4."
)


DONORS = ["donor 1", "donor 2", "donor 3", "donor 4"]

# The three levels every ratio was tested across, and the three Dunn comparisons Prism ran
# all-vs-all. 
CONDITIONS = ["D2", "D8A", "D8C"]
COMPARISONS = [("D2", "D8A"), ("D2", "D8C"), ("D8A", "D8C")]

# The suffix the collaborator put on every sheet but GSSG:GSH. The normalization belongs in
# the row-2 methods sentence, not repeated in every row label.
NORMALIZED_SUFFIX = " (normalized to D2)"

# Row order, by ratio name after the suffix is stripped: the order the ratios appear in the
# Prism document's data-sheet list. 
RATIO_ORDER = [
    "Adenosine:AMP",
    "Guanosine:GMP",
    "Uridine:UMP",
    "ATP:ADP",
    "GTP:GDP",
    "UTP:UDP",
    "GSSG:GSH",
]

# GTP:GDP has one fewer donor because one exp didn't pass metabolite QC
EXPECTED_UNMATCHED = {("GTP:GDP", 3)}


def read_json(archive, name):
    return json.loads(archive.read(name))


def read_table(archive, table_uid):
    """The rows of a data table as floats, blank cells as None."""
    text = archive.read(f"data/tables/{table_uid}/data.csv").decode()
    rows = []
    for row in csv.reader(io.StringIO(text)):
        if not any(cell.strip() for cell in row):
            continue  # Prism pads a table out to its allocated row count.
        rows.append([float(cell) if cell.strip() else None for cell in row])
    return rows


def dunn_p_values(multiple_comparisons):
    """The Dunn adjusted p-values as {(level_a, level_b): p}.
    """
    p_values = {}
    for outer in multiple_comparisons["tests"].values():
        a = outer["id"]["col"]["title"]
        for inner in outer["value"].values():
            b = inner["id"]["col"]["title"]
            p_values[(a, b)] = inner["value"]["pValueAdjusted"]["value"]
    return p_values


def extract(prism_path):
    """One record per Friedman analysis: name, per-condition values, means, n, p-values.

    Returns records in `RATIO_ORDER`.
    """
    records = []
    with zipfile.ZipFile(prism_path) as archive:
        document = read_json(archive, "document.json")
        titles = document["sheetAttributesMap"]

        for uid in document["sheets"]["analyses"]:
            sheet = read_json(archive, f"analyses/{uid}/sheet.json")
            results = read_json(archive, f"analyses/{uid}/results.json")["content"]
            model = results["friedmanTestModel"]

            data_uid = sheet["inputSheets"][0]["uid"]
            data_sheet = read_json(archive, f"data/sheets/{data_uid}/sheet.json")
            rows = read_table(archive, data_sheet["table"]["uid"])

            name = titles[data_uid]["title"].removesuffix(NORMALIZED_SUFFIX)

            # The column each level was actually read from. Selecting by this index rather
            # than by position is what makes GSSG:GSH's ten-column table come out right.
            columns = {
                stat["id"]["col"]["title"]: stat["id"]["col"]["idx"]
                for stat in results["descriptiveStatistics"].values()
            }
            means = {
                stat["id"]["col"]["title"]: stat["value"]["mean"]
                for stat in results["descriptiveStatistics"].values()
            }
            assert sorted(columns) == sorted(CONDITIONS), f"{name}: levels {sorted(columns)}"

            values = {
                condition: [row[columns[condition]] for row in rows]
                for condition in CONDITIONS
            }
            for condition, donors in values.items():
                assert len(donors) == len(DONORS), (
                    f"{name} {condition}: {len(donors)} rows, expected {len(DONORS)}"
                )

            p_values = dunn_p_values(model["multipleComparisons"])
            assert sorted(p_values) == sorted(COMPARISONS), (
                f"{name}: comparisons {sorted(p_values)}"
            )

            records.append(
                {
                    "name": name,
                    "values": values,
                    "means": means,
                    "n": model["dataSummary"]["matchedCount"],
                    "friedman_p": model["friedmanTest"]["pValue"]["value"],
                    "dunn_p": p_values,
                }
            )

    found = {record["name"] for record in records}
    assert found == set(RATIO_ORDER), f"ratios in file: {sorted(found)}"
    by_name = {record["name"]: record for record in records}
    return [by_name[name] for name in RATIO_ORDER]


def columns():
    """Header row: donor columns per condition, then the means, n and the four p-values."""
    header = [f"{condition} {donor}" for condition in CONDITIONS for donor in DONORS]
    header = ["ratio"] + header
    header += [f"mean {condition}" for condition in CONDITIONS]
    header += ["n donors", "Friedman p"]
    header += [f"Dunn adj. p, {a} vs {b}" for a, b in COMPARISONS]
    return header


def build_rows(records):
    """The table body: only the donors the Friedman test kept, and its statistics.

    A donor Friedman excluded is blank in every condition, not just the one it was missing.
    That keeps each row self-consistent - the mean of the values shown IS the mean reported,
    and `n donors` counts exactly the donors with values - rather than listing measurements
    that no statistic in the row was computed from.
    """
    rows, unmatched = [], set()
    for record in records:
        matched = matched_donors(record)
        unmatched |= {
            (record["name"], index)
            for index in range(len(DONORS))
            if index not in matched
        }

        row = [record["name"]]
        for condition in CONDITIONS:
            for index in range(len(DONORS)):
                row.append(record["values"][condition][index] if index in matched else None)
        row += [record["means"][condition] for condition in CONDITIONS]
        row += [record["n"], record["friedman_p"]]
        row += [record["dunn_p"][pair] for pair in COMPARISONS]
        rows.append(row)

    assert unmatched == EXPECTED_UNMATCHED, f"unmatched donors are {sorted(unmatched)}"
    return rows


def matched_donors(record):
    """The donor indices with a value in every condition - the rows the Friedman test used.

    Friedman is a repeated-measures test, so Prism drops a donor from the whole analysis if
    it is missing any one condition. That is why GTP:GDP reports n = 3: its fourth donor has
    no D8A value, so that donor is out of all three columns, not just D8A.
    """
    return [
        index
        for index in range(len(DONORS))
        if all(record["values"][condition][index] is not None for condition in CONDITIONS)
    ]


def check_means(records):
    """Every mean Prism reports must be the matched-case mean of the values pulled for it.

    This is the cross-check that the columns were not misaligned: a mean taken from the
    results and values taken from the wrong table column would not agree. It doubles as a
    check that `matched_donors` reproduces Prism's own case-wise deletion, since a
    disagreement about which donors are in would move the mean.
    """
    for record in records:
        matched = matched_donors(record)
        assert len(matched) == record["n"], (
            f"{record['name']}: {len(matched)} matched donors, Prism reports n = "
            f"{record['n']}"
        )
        for condition in CONDITIONS:
            observed = [record["values"][condition][index] for index in matched]
            mean = sum(observed) / len(observed)
            assert abs(mean - record["means"][condition]) < 1e-6, (
                f"{record['name']} {condition}: computed {mean}, "
                f"Prism reports {record['means'][condition]}"
            )


def build_sheet(rows):
    """The workbook in house style: title, description, header on row 3, striped table."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = S3_2_SHEET
    worksheet["A1"] = S3_2_TITLE
    worksheet["A2"] = S3_2_DESCRIPTION
    # Writing A2 leaves openpyxl's append cursor on row 2, so the header lands on row 3 - the
    # same layout as every other sheet, where the two title rows precede the header.
    worksheet.append(columns())
    for row in rows:
        worksheet.append(row)

    last_col = worksheet.cell(row=3, column=len(columns())).column_letter
    worksheet.add_table(
        Table(
            displayName="Frame08",
            ref=f"A3:{last_col}{3 + len(rows)}",
            tableStyleInfo=TableStyleInfo(showRowStripes=True),
        )
    )
    return workbook


def verify(path, rows):
    """Reopen the written workbook and check the layout and every cell of the body."""
    worksheet = load_workbook(path)[S3_2_SHEET]
    assert worksheet["A1"].value.startswith("S3-2 "), worksheet["A1"].value
    assert worksheet["A2"].value == S3_2_DESCRIPTION
    header = [cell.value for cell in worksheet[3]]
    assert header == columns(), header

    for index, expected in enumerate(rows):
        written = [cell.value for cell in worksheet[4 + index]]
        assert len(written) == len(expected), f"{expected[0]}: {len(written)} cells"
        for name, got, want in zip(header, written, expected):
            same = (
                math.isclose(got, want, rel_tol=1e-12)
                if isinstance(want, (int, float))
                else got == want
            )
            assert same, f"{expected[0]}, {name}: wrote {want!r}, read back {got!r}"
    assert worksheet.max_row == 3 + len(rows), worksheet.max_row
    print(f"  verified: {len(rows)} ratios, {len(header)} columns, header on row 3")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("prism", nargs="?", type=Path, default=PRISM)
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()

    records = extract(args.prism)
    check_means(records)
    rows = build_rows(records)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    build_sheet(rows).save(args.out)
    print(f"wrote {args.out}")
    verify(args.out, rows)

    for record in records:
        print(
            f"  {record['name']:16s} n={record['n']}  "
            f"Friedman p={record['friedman_p']:.4f}  "
            + "  ".join(
                f"{a} vs {b}={record['dunn_p'][(a, b)]:.4f}" for a, b in COMPARISONS
            )
        )

    # The `contents` row for Data S3.xlsx, in the #/Title/Description shape that sheet uses.
    print("\ncontents entry:")
    print(f"  2\t{S3_2_SHEET}\t{S3_2_TITLE.removeprefix('S3-2 ')}")


if __name__ == "__main__":
    main()
